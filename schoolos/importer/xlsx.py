"""Build the import template and read uploaded sheets (openpyxl only, no Frappe imports)."""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from schoolos.importer.parse import COLUMNS, map_headers

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
REQ_FILL = PatternFill("solid", fgColor="B42318")
EXAMPLE_FONT = Font(name="Arial", italic=True, color="808080")


def build_template(school_name="", academic_year=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    for i, (key, header, required, example, help_text) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=header + (" *" if required else ""))
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = REQ_FILL if required else HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if help_text:
            c.comment = Comment(help_text, "SchoolOS")
        e = ws.cell(row=2, column=i, value=example)
        e.font = EXAMPLE_FONT
        ws.column_dimensions[c.column_letter].width = max(12, len(header) + 4)
        if key in ("date_of_birth", "joining_date", "student_mobile", "father_mobile", "mother_mobile",
                   "pincode", "admission_no"):
            for r in range(2, 3002):
                ws.cell(row=r, column=i).number_format = "@"  # keep as text (no 9.87E+09)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "E2"

    col = {k: i for i, (k, *_r) in enumerate(COLUMNS, 1)}
    gender = DataValidation(type="list", formula1='"Male,Female,Other"', allow_blank=False)
    blood = DataValidation(type="list", formula1='"A+,A-,B+,B-,AB+,AB-,O+,O-"', allow_blank=True)
    ws.add_data_validation(gender)
    ws.add_data_validation(blood)
    gl = ws.cell(row=1, column=col["gender"]).column_letter
    bl = ws.cell(row=1, column=col["blood_group"]).column_letter
    gender.add(f"{gl}2:{gl}3001")
    blood.add(f"{bl}2:{bl}3001")

    info = wb.create_sheet("How to fill")
    lines = [
        f"Student import for {school_name or 'your school'}" + (f" · {academic_year}" if academic_year else ""),
        "",
        "1. One row per student on the 'Students' sheet. Row 2 is an example: replace or delete it.",
        "2. Red columns (*) are required: Class, Section, First Name, Gender, Date of Birth.",
        "3. Dates as DD-MM-YYYY (e.g. 14-05-2016).",
        "4. Mobiles as 10 digits (e.g. 9876543210). These receive WhatsApp alerts.",
        "5. Brothers/sisters: use the same parent mobile; they'll be linked to one parent account.",
        "6. Admission / GR No lets you re-upload the sheet later to update students instead of duplicating them.",
        "7. Upload in SchoolOS Student Import, click Validate, fix any red rows, then Start Import.",
        "",
        "Columns:",
    ] + [f"   {h}{' (required)' if r else ''}: {t}" for _, h, r, _, t in COLUMNS]
    for i, line in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=line)
        c.font = Font(name="Arial", bold=(i == 1), size=13 if i == 1 else 11)
    info.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(content):
    """Read the first sheet (or 'Students'). Returns (rows as dicts keyed by column key, missing_headers)."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb["Students"] if "Students" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None) or []
    mapping = map_headers(header)
    found = set(mapping.values())
    missing = [h for k, h, req, *_ in COLUMNS if req and k not in found]
    rows = []
    for values in it:
        rows.append({key: (values[i] if i < len(values) else None) for i, key in mapping.items()})
    wb.close()
    return rows, missing


def is_example_row(raw):
    """True if the row is the untouched example from the template."""
    ex = {k: e for k, _, _, e, _ in COLUMNS}
    return str(raw.get("first_name") or "") == ex["first_name"] and str(raw.get("admission_no") or "") == ex["admission_no"]
