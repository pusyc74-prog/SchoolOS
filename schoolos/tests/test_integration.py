"""End-to-end tests on a real Frappe site, against the demo school.

    bench --site <test site> set-config allow_tests true
    bench --site <test site> run-tests --app schoolos

The demo (schoolos.demo.make_demo) is created on first run if it isn't there. Never run this on a
live school's site: Frappe refuses unless allow_tests is set, and the demo adds fake students.
"""

import io
import xml.etree.ElementTree as ET

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import add_days, getdate, nowdate, today

from schoolos.demo import GROUP, PROGRAM

PRINCIPAL = "principal@demo.school"
TEACHER = "teacher@demo.school"
PARENT = "parent@demo.school"
OTHER_PARENT_CHILD_INDEX = 5  # a child that isn't the demo parent's


def ensure_demo():
    if not frappe.db.exists("Student Group", GROUP):
        from schoolos.demo import make_demo

        make_demo(strict=1)
        frappe.db.commit()


def group_students():
    return frappe.get_all("Student Group Student", filters={"parent": GROUP, "active": 1},
                          order_by="idx asc", pluck="student")


def parent_children():
    from schoolos.api.common import guardian_children

    with IntegrationTestCase.set_user(PARENT):
        return [c.name for c in guardian_children()]


def is_holiday(day):
    return bool(frappe.db.exists("Holiday", {"holiday_date": day})) or getdate(day).weekday() == 6


def school_day(offset):
    """The first school day at least `offset` days from today (skips Sundays and holidays)."""
    d = add_days(today(), offset)
    while is_holiday(d):
        d = add_days(d, 1)
    return d


def take_response():
    """Read and clear a file response set by an endpoint (PDF / XML / XLSX)."""
    r = frappe.local.response
    out = (r.get("type"), r.get("filename"), r.get("filecontent"))
    for k in ("type", "filename", "filecontent"):
        r.pop(k, None)
    return out


class SchoolOSTestCase(IntegrationTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        ensure_demo()
        cls.students = group_students()
        cls.children = parent_children()
        cls.other_child = next(s for s in cls.students if s not in cls.children)

    def setUp(self):
        frappe.local.response = frappe._dict()


# ---------------------------------------------------------------- install
class TestInstall(SchoolOSTestCase):
    def test_roles_fields_scales(self):
        self.assertTrue(frappe.db.exists("Role", "School Principal"))
        for dt, field in [("Student", "schoolos_admission_no"), ("Guardian", "schoolos_alerts_opt_in"),
                          ("Student Leave Application", "schoolos_status"), ("Purchase Invoice", "schoolos_rejected"),
                          ("Student Log", "schoolos_share_with_parent")]:
            self.assertTrue(frappe.get_meta(dt).has_field(field), f"{dt}.{field} missing")
        for scale in ("CBSE 8-Point", "Maharashtra 9-Point"):
            self.assertEqual(frappe.db.get_value("Grading Scale", scale, "docstatus"), 1)
        self.assertTrue(frappe.db.exists("Print Format", "SchoolOS Report Card"))

    def test_migrate_is_idempotent(self):
        from schoolos.install import setup

        setup()
        setup()  # second run must not fail or duplicate

    def test_all_doctypes_load(self):
        for dt in frappe.get_all("DocType", filters={"module": "SchoolOS"}, pluck="name"):
            frappe.get_meta(dt)
            if not frappe.get_meta(dt).issingle and not frappe.get_meta(dt).istable:
                frappe.get_all(dt, limit=1)

    def test_demo_is_complete(self):
        self.assertEqual(len(self.students), 30)
        self.assertEqual(len(self.children), 2)
        self.assertGreater(frappe.db.count("Student Attendance", {"student_group": GROUP, "docstatus": 1}), 300)
        self.assertGreater(frappe.db.count("Sales Invoice", {"docstatus": 1, "student": ["is", "set"]}), 30,
                           "demo fee invoices were not created")
        self.assertGreater(frappe.db.count("Payment Entry", {"docstatus": 1}), 10)
        self.assertGreater(frappe.db.count("SchoolOS Lesson Log"), 10)
        self.assertGreater(frappe.db.count("Student Report Card", {"docstatus": 1}), 0)
        self.assertTrue(frappe.db.exists("SchoolOS Album", {"published": 1}))


# ---------------------------------------------------------------- dashboards & access
class TestDashboards(SchoolOSTestCase):
    def test_context_per_role(self):
        from schoolos.api.dashboard import get_context

        with self.set_user(PRINCIPAL):
            v = get_context()["views"]
            for view in ("principal", "fees", "accounts", "calendar"):
                self.assertIn(view, v)
        with self.set_user(TEACHER):
            v = get_context()["views"]
            self.assertIn("teacher", v)
            self.assertNotIn("principal", v)
            self.assertNotIn("fees", v)
        with self.set_user(PARENT):
            ctx = get_context()
            self.assertIn("parent", ctx["views"])
            self.assertNotIn("principal", ctx["views"])
            self.assertTrue(ctx["branding"]["school_name"])

    def test_principal(self):
        from schoolos.api.dashboard import principal

        with self.set_user(PRINCIPAL):
            d = principal()
        self.assertIsInstance(d, dict)
        self.assertTrue(d)

    def test_teacher(self):
        from schoolos.api.dashboard import teacher

        with self.set_user(TEACHER):
            d = teacher()
        self.assertTrue(d["instructor"])
        self.assertIn(GROUP, [g.get("student_group") or g.get("name") for g in d["groups_today"]])

    def test_parent(self):
        from schoolos.api.dashboard import parent

        with self.set_user(PARENT):
            d = parent()
        self.assertEqual(len(d["children"]), 2)
        c = d["children"][0]
        self.assertIsNotNone(c["attendance"])
        self.assertIsNotNone(c["fees"], "fees brief missing for parent")
        self.assertTrue(c["calendar"])

    def test_student_view_and_privacy(self):
        from schoolos.api.dashboard import student

        with self.set_user(PARENT):
            d = student(self.children[0])
            self.assertEqual(d["student"]["name"], self.children[0])
            with self.assertRaises(frappe.PermissionError):
                student(self.other_child)

    def test_role_walls(self):
        from schoolos.api.accounts import summary as acc
        from schoolos.api.dashboard import principal, teacher
        from schoolos.api.fees import summary as fees

        with self.set_user(PARENT):
            for fn in (principal, fees, acc, teacher):
                with self.assertRaises(frappe.PermissionError, msg=fn.__module__ + "." + fn.__name__):
                    fn()
        with self.set_user(TEACHER):
            for fn in (principal, fees, acc):
                with self.assertRaises(frappe.PermissionError, msg=fn.__module__ + "." + fn.__name__):
                    fn()
        with self.set_user("Guest"):
            from schoolos.api.dashboard import parent

            with self.assertRaises(frappe.PermissionError):
                parent()


# ---------------------------------------------------------------- module switches
class TestModules(SchoolOSTestCase):
    def test_switch_off_hides_and_blocks(self):
        from schoolos.api.dashboard import get_context
        from schoolos.api.fees import for_student, summary
        from schoolos.modules import MODULES

        for key in MODULES:
            frappe.db.set_single_value("SchoolOS Settings", f"module_{key}", 0)
        frappe.clear_document_cache("SchoolOS Settings", "SchoolOS Settings")
        try:
            with self.set_user(PRINCIPAL):
                ctx = get_context()
                self.assertNotIn("fees", ctx["views"])
                self.assertNotIn("accounts", ctx["views"])
                with self.assertRaises(frappe.PermissionError):
                    summary()
                from schoolos.api.dashboard import principal

                principal()  # must still render with every module off
            with self.set_user(TEACHER):
                from schoolos.api.dashboard import teacher

                teacher()
            with self.set_user(PARENT):
                with self.assertRaises(frappe.PermissionError):
                    for_student(self.children[0])
                from schoolos.api.dashboard import parent

                d = parent()  # still works, just without the switched-off parts
                self.assertIsNone(d["children"][0]["fees"])
                from schoolos.api.dashboard import student

                student(self.children[0])
        finally:
            for key in MODULES:
                frappe.db.set_single_value("SchoolOS Settings", f"module_{key}", 1)
            frappe.clear_document_cache("SchoolOS Settings", "SchoolOS Settings")


# ---------------------------------------------------------------- fees
class TestFees(SchoolOSTestCase):
    def test_summary(self):
        from schoolos.api.fees import summary

        with self.set_user(PRINCIPAL):
            s = summary()
        self.assertGreater(s["billed"], 0)
        self.assertGreater(s["collected"], 0)
        self.assertGreater(s["outstanding"], 0)
        self.assertTrue(s["defaulters"])
        self.assertEqual(len(s["ageing"]), 4)
        self.assertAlmostEqual(s["billed"], s["collected"] + s["outstanding"], places=1)

    def test_parent_fees_and_receipt(self):
        from schoolos.api.fees import for_student, receipt

        with self.set_user(PARENT):
            f = for_student(self.children[0])
            self.assertTrue(f["invoices"])
            paid = [i for i in f["invoices"] if i["outstanding"] < i["total"]]
            if paid:
                receipt(paid[0]["invoice"])
                kind, _name, content = take_response()
                self.assertEqual(kind, "pdf")
                self.assertEqual(bytes(content[:4]), b"%PDF")
            with self.assertRaises(frappe.PermissionError):
                for_student(self.other_child)

    def test_reminders_dry_run(self):
        from schoolos.api.fees import remind, send_scheduled_reminders

        before = frappe.db.count("SchoolOS Alert Log", {"alert_type": "Fee Reminder"})
        with self.set_user(PRINCIPAL):
            r = remind()
        self.assertGreater(r["students"], 0)
        self.assertGreater(frappe.db.count("SchoolOS Alert Log", {"alert_type": "Fee Reminder"}), before)
        send_scheduled_reminders()  # must not raise


# ---------------------------------------------------------------- alerts
class TestAlerts(SchoolOSTestCase):
    def test_deliver_dry_run(self):
        from schoolos.alerts.engine import deliver, retry_failed

        self.assertEqual(frappe.db.get_single_value("SchoolOS Settings", "dry_run"), 1)
        log = frappe.get_all("SchoolOS Alert Log", filters={"status": "Queued"}, limit=1, pluck="name")
        if log:
            deliver(log[0])
            self.assertEqual(frappe.db.get_value("SchoolOS Alert Log", log[0], "status"), "Dry Run")
        retry_failed()

    def test_absence_alert_on_attendance(self):
        if is_holiday(today()):
            self.skipTest("today is a holiday; attendance can't be marked")
        s = self.students[-1]
        frappe.db.delete("Student Attendance", {"student": s, "date": today()})
        # (alert logs are not deleted: the background worker may be sending them right now;
        #  one absence alert per child per day is deduplicated, so "exists" is the right check)
        doc = frappe.get_doc({"doctype": "Student Attendance", "student": s, "student_group": GROUP,
                              "date": today(), "status": "Absent"})
        doc.insert()
        doc.submit()
        self.assertTrue(frappe.db.exists("SchoolOS Alert Log", {"alert_type": "Absence", "student": s}),
                        "absence alert not queued")

    def test_send_test_message(self):
        from schoolos.alerts.engine import send_test

        rows = send_test(email="office@example.com")
        self.assertTrue(rows)
        self.assertTrue(all(r.status in ("Dry Run", "Queued") for r in rows))


# ---------------------------------------------------------------- homework & learning
class TestHomeworkLearning(SchoolOSTestCase):
    def test_post_and_see_homework(self):
        from schoolos.api.homework import for_student, post, post_options

        with self.set_user(TEACHER):
            opts = post_options()
            self.assertIn(GROUP, [g.get("name") if isinstance(g, dict) else g for g in opts["groups"]])
            r = post(GROUP, "Integration test homework", "Homework", None, "Do exercise 4.2",
                     school_day(2))
            self.assertEqual(r["students"], 30)
        with self.set_user(PARENT):
            titles = [h.title for h in for_student(self.children[0])]
            self.assertIn("Integration test homework", titles)

    def test_teacher_cannot_post_elsewhere(self):
        from schoolos.api.homework import post

        other = frappe.get_doc({"doctype": "Student Group", "student_group_name": "Other Group IT",
                                "group_based_on": "Activity", "academic_year": frappe.db.get_value("Student Group", GROUP, "academic_year")})
        other.insert(ignore_permissions=True)
        with self.set_user(TEACHER):
            with self.assertRaises(frappe.PermissionError):
                post(other.name, "Not mine", "Homework", None, "x", school_day(1))
            from schoolos.api.learning import group_students, post_note

            with self.assertRaises(frappe.PermissionError):
                group_students(other.name)
        imported = frappe.get_doc({"doctype": "Student", "first_name": "Outsider", "gender": "Male",
                                   "student_email_id": "outsider@students.noemail",
                                   "date_of_birth": "2015-01-01"}).insert(ignore_permissions=True)
        with self.set_user(TEACHER):
            from schoolos.api.learning import post_note

            with self.assertRaises(frappe.PermissionError):
                post_note(imported.name, "not my student")

    def test_progress_activity(self):
        from schoolos.api.learning import activity, student_progress, syllabus_overview

        with self.set_user(PARENT):
            p = student_progress(self.children[0])
            self.assertTrue(p["subjects"])
            maths = next(s for s in p["subjects"] if "Math" in s["course"])
            self.assertIsNotNone(maths["syllabus"], "Mathematics has no syllabus")
            self.assertIn(maths["syllabus"]["status"], ("behind", "on track", "ahead"))
            self.assertTrue(activity(self.children[0]))
        with self.set_user(PRINCIPAL):
            o = syllabus_overview()
            self.assertTrue(o["rows"])

    def test_log_lesson_and_note(self):
        from schoolos.api.learning import lesson_options, post_lesson, post_note

        with self.set_user(TEACHER):
            opts = lesson_options()
            self.assertIn(GROUP, opts["groups"])
            subject = opts["courses"][GROUP][0]
            todo = [t for t in subject["topics"] if t["state"] == "todo"]
            topic = todo[0]["topic"] if todo else None
            r = post_lesson(GROUP, subject["course"], topic)
            self.assertTrue(r["name"])
            if topic:
                self.assertEqual(r["coverage"]["covered"], subject.get("covered", r["coverage"]["covered"] - 1) + 1)
            n = post_note(self.children[0], "Integration test note", "General", 1)
            self.assertTrue(n["name"])


# ---------------------------------------------------------------- leave
class TestLeave(SchoolOSTestCase):
    def test_apply_approve_marks_leave(self):
        from schoolos.api.leave import apply, decide, for_student, pending

        child = self.children[1]
        day = school_day(5)
        with self.set_user(PARENT):
            r = apply(child, day, day, "Doctor appointment")
            self.assertEqual(r["status"], "Pending")
            self.assertIn(r["name"], [x["name"] for x in for_student(child)])
            with self.assertRaises(frappe.ValidationError):
                apply(child, day, add_days(day, -1), "bad dates")
            with self.assertRaises(frappe.PermissionError):
                apply(self.other_child, day, day, "not my child")
        with self.set_user(TEACHER):
            self.assertIn(r["name"], [x["name"] for x in pending()])
            out = decide(r["name"], 1, "Get well soon")
        self.assertEqual(out["status"], "Approved")
        self.assertEqual(frappe.db.get_value("Student Leave Application", r["name"], "docstatus"), 1)
        self.assertEqual(frappe.db.get_value("Student Attendance",
                                                 {"student": child, "date": day, "docstatus": 1}, "status"), "Leave")

    def test_reject_needs_note_and_withdraw(self):
        from schoolos.api.leave import apply, decide, withdraw

        child = self.children[0]
        with self.set_user(PARENT):
            d1, d2 = school_day(8), school_day(12)
            r1 = apply(child, d1, d1, "Wedding")
            r2 = apply(child, d2, d2, "Trip")
            self.assertEqual(withdraw(r2["name"]), "ok")
        with self.set_user(TEACHER):
            with self.assertRaises(frappe.ValidationError):
                decide(r1["name"], 0, "")
            out = decide(r1["name"], 0, "Unit test that day")
        self.assertEqual(out["status"], "Rejected")
        self.assertEqual(frappe.db.get_value("Student Leave Application", r1["name"], "docstatus"), 0)

    def test_approved_leave_suppresses_absence_alert(self):
        from schoolos.api.leave import apply, decide

        if is_holiday(today()):
            self.skipTest("today is a holiday")

        child = self.children[1]
        day = today()
        frappe.db.delete("Student Attendance", {"student": child, "date": day})
        before = frappe.db.count("SchoolOS Alert Log", {"alert_type": "Absence", "student": child,
                                                        "creation": [">=", today()]})
        with self.set_user(PARENT):
            r = apply(child, day, day, "Fever")
        with self.set_user(TEACHER):
            decide(r["name"], 1, "")
        self.assertEqual(frappe.db.get_value("Student Attendance", {"student": child, "date": day, "docstatus": 1},
                                             "status"), "Leave")
        self.assertEqual(frappe.db.count("SchoolOS Alert Log", {"alert_type": "Absence", "student": child,
                                                                 "creation": [">=", today()]}), before)


# ---------------------------------------------------------------- calendar & gallery
class TestCalendar(SchoolOSTestCase):
    def test_month_and_albums(self):
        from schoolos.api.school_calendar import album, albums, month

        t = getdate(today())
        with self.set_user(PARENT):
            m = month(t.year, t.month, self.children[0])
            self.assertTrue(m["upcoming"], "no upcoming events for parent")
            a = albums(self.children[0])
            self.assertTrue(a)
            full = album(a[0]["name"])
            self.assertEqual(len(full["photos"]), 5)
        with self.set_user(PRINCIPAL):
            self.assertTrue(month(t.year, t.month)["items"] is not None)


# ---------------------------------------------------------------- report cards
class TestReportCards(SchoolOSTestCase):
    def test_generate_and_download(self):
        from schoolos.api.report_card import download, generate

        template = frappe.get_all("Report Card Template", filters={"program": PROGRAM}, pluck="name")[0]
        with self.set_user(PRINCIPAL):
            r = generate(template, GROUP)
        self.assertEqual(r["created"] + r["refreshed"] + r["already_published"], 30)

        card = frappe.get_all("Student Report Card", filters={"docstatus": 1, "student": ["in", self.children]},
                              pluck="name")[0]
        doc = frappe.get_doc("Student Report Card", card)
        self.assertTrue(doc.subjects)
        self.assertGreater(doc.overall_percentage or 0, 0)
        with self.set_user(PARENT):
            download(card)
            kind, _name, content = take_response()
            self.assertEqual(kind, "pdf")
            self.assertEqual(bytes(content[:4]), b"%PDF")
        other = frappe.get_all("Student Report Card", filters={"docstatus": 1, "student": ["not in", self.children]},
                               pluck="name")
        if other:
            with self.set_user(PARENT), self.assertRaises(frappe.PermissionError):
                download(other[0])


# ---------------------------------------------------------------- accounts & Tally
class TestAccounts(SchoolOSTestCase):
    def test_full_mode_summary(self):
        from schoolos.api.accounts import summary

        frappe.db.set_single_value("SchoolOS Settings", "accounts_mode", "Full accounting in SchoolOS")
        frappe.clear_document_cache("SchoolOS Settings", "SchoolOS Settings")
        with self.set_user(PRINCIPAL):
            s = summary()
        self.assertEqual(s["mode"], "full")
        self.assertGreater(s["year"]["income"], 0, "fee income not visible in accounts")
        self.assertTrue(s["monthly"])
        self.assertIsInstance(s["approvals"], list)
        self.assertIn("balances", s)

    def test_approve_and_reject_bills(self):
        from schoolos.api.accounts import _approvals, _company, decide

        with self.set_user(PRINCIPAL):
            bills = [b for b in _approvals(_company()) if b["doctype"] == "Purchase Invoice"]
            self.assertGreaterEqual(len(bills), 2)
            ok = decide("Purchase Invoice", bills[0]["name"], 1)
            self.assertEqual(ok["status"], "Approved")
            self.assertEqual(frappe.db.get_value("Purchase Invoice", bills[0]["name"], "docstatus"), 1)
            self.assertTrue(frappe.db.exists("Comment", {"reference_name": bills[0]["name"],
                                                         "content": ["like", f"%{PRINCIPAL}%"]}))
            self.assertEqual(frappe.session.user, PRINCIPAL)  # user restored after posting
            with self.assertRaises(frappe.ValidationError):
                decide("Purchase Invoice", bills[1]["name"], 0, "")
            decide("Purchase Invoice", bills[1]["name"], 0, "Wrong amount")
            self.assertEqual(frappe.db.get_value("Purchase Invoice", bills[1]["name"], "schoolos_rejected"), 1)
            with self.assertRaises(frappe.ValidationError):
                decide("Purchase Invoice", bills[0]["name"], 1)  # already processed
        with self.set_user(TEACHER), self.assertRaises(frappe.PermissionError):
            decide("Purchase Invoice", bills[1]["name"], 1)

    def test_tally_exports(self):
        from schoolos.api.accounts import tally_export, tally_preview

        frappe.db.set_single_value("SchoolOS Settings", "accounts_mode", "Tally (export fees to Tally)")
        frappe.clear_document_cache("SchoolOS Settings", "SchoolOS Settings")
        try:
            ys = frappe.db.get_value("Academic Year", frappe.db.get_value("Student Group", GROUP, "academic_year"),
                                     "year_start_date")
            f, t = str(ys), nowdate()
            with self.set_user(PRINCIPAL):
                p = tally_preview(f, t)
                self.assertGreater(p["invoices"], 0)
                self.assertGreater(p["receipts"], 0)

                tally_export(f, t, "masters", "xml")
                _k, name, xml = take_response()
                root = ET.fromstring(xml)
                self.assertTrue(root.findall(".//LEDGER"), "no ledgers in masters file")

                tally_export(f, t, "vouchers", "xml")
                _k, name, xml = take_response()
                root = ET.fromstring(xml)
                vouchers = root.findall(".//VOUCHER")
                self.assertEqual(len(vouchers), p["invoices"] + p["receipts"])
                for v in vouchers:  # every voucher balances
                    amounts = [float(a.text) for a in v.findall(".//AMOUNT")]
                    self.assertAlmostEqual(sum(amounts), 0, places=2)

                tally_export(f, t, "vouchers", "xlsx")
                _k, name, content = take_response()
                from openpyxl import load_workbook

                self.assertGreater(load_workbook(io.BytesIO(content)).active.max_row, 1)
            self.assertEqual(str(frappe.db.get_single_value("SchoolOS Settings", "tally_last_export_on")), nowdate())
        finally:
            frappe.db.set_single_value("SchoolOS Settings", "accounts_mode", "Full accounting in SchoolOS")
            frappe.clear_document_cache("SchoolOS Settings", "SchoolOS Settings")


# ---------------------------------------------------------------- import
class TestImport(SchoolOSTestCase):
    def test_excel_import_and_reimport(self):
        from openpyxl import Workbook

        from schoolos.importer.parse import COLUMNS
        from schoolos.importer.run import run

        year = frappe.db.get_value("Student Group", GROUP, "academic_year")
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append([c[1] for c in COLUMNS])
        rows = [
            {"Class": "Class 6", "Section": "B", "Roll No": 1, "Admission / GR No": "IT-001", "First Name": "Test",
             "Last Name": "Kale", "Gender": "F", "Date of Birth": "01-02-2015", "Father Name": "Ravi Kale",
             "Father Mobile": "9800000001"},
            {"Class": "Class 6", "Section": "B", "Roll No": 2, "Admission / GR No": "IT-002", "First Name": "Sibling",
             "Last Name": "Kale", "Gender": "M", "Date of Birth": "03-04-2016", "Father Name": "Ravi Kale",
             "Father Mobile": "+91 98000 00001"},
            {"Class": "Class 6", "Section": "B", "First Name": "Bad", "Gender": "X", "Date of Birth": "31-02-2015"},
        ]
        for r in rows:
            ws.append([r.get(c[1]) for c in COLUMNS])
        buf = io.BytesIO()
        wb.save(buf)

        def do_import():
            imp = frappe.get_doc({"doctype": "SchoolOS Student Import", "academic_year": year,
                                  "create_missing_classes": 1}).insert()
            f = frappe.get_doc({"doctype": "File", "file_name": "it-import.xlsx", "content": buf.getvalue(),
                                "attached_to_doctype": imp.doctype, "attached_to_name": imp.name, "is_private": 1}).insert()
            imp.import_file = f.file_url
            imp.save()
            res = imp.check_file()
            self.assertEqual(res["ready"], 2)
            run(imp.name)
            imp.reload()
            return imp

        imp = do_import()
        self.assertEqual(imp.imported_rows, 2, imp.log)
        self.assertEqual(imp.status, "Completed with Errors")
        s1 = frappe.db.get_value("Student", {"schoolos_admission_no": "IT-001"})
        s2 = frappe.db.get_value("Student", {"schoolos_admission_no": "IT-002"})
        self.assertTrue(s1 and s2)
        g1 = frappe.get_all("Student Guardian", filters={"parent": s1}, pluck="guardian")
        g2 = frappe.get_all("Student Guardian", filters={"parent": s2}, pluck="guardian")
        self.assertTrue(set(g1) & set(g2), "siblings should share the father's Guardian record")
        self.assertTrue(frappe.db.exists("Student Group", {"program": "Class 6"}))

        count = frappe.db.count("Student")
        do_import()  # re-upload updates, never duplicates
        self.assertEqual(frappe.db.count("Student"), count)


# ---------------------------------------------------------------- PWA & page
class TestPWA(SchoolOSTestCase):
    def test_manifest_and_worker(self):
        import json

        from schoolos.api.pwa import manifest, service_worker

        m = json.loads(manifest().get_data())
        self.assertTrue(m["name"])
        self.assertEqual(m["start_url"][:7], "/school")
        self.assertIn(b"fetch", service_worker().get_data())

    def test_school_page_renders(self):
        from frappe.utils import set_request
        from frappe.website.serve import get_response_content

        with self.set_user(PARENT):
            set_request(method="GET", path="/school")
            html = get_response_content("school")
        self.assertIn("school.js", html)
