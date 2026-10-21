"""Importer and message tests that don't need a Frappe site.  Run: python3 -m pytest tests -q"""

from datetime import date, datetime

from openpyxl import load_workbook

from schoolos.alerts.messages import email_template, localise_context, render, template_params
from schoolos.importer.parse import COLUMNS, check_rows, clean_row, group_name, map_headers, parse_date
from schoolos.importer.xlsx import build_template, is_example_row, read_rows

TODAY = date(2026, 10, 21)


def raw(**kw):
    base = {"program": "Class 5", "section": "A", "first_name": "Riya", "last_name": "Patil",
            "gender": "F", "date_of_birth": "14-05-2016", "father_mobile": "98765 43210"}
    base.update(kw)
    return base


def test_headers_tolerant():
    m = map_headers(["Class *", " section ", "FIRST NAME", "Unknown", "Father  Mobile"])
    assert m == {0: "program", 1: "section", 2: "first_name", 4: "father_mobile"}


def test_dates():
    assert parse_date("14-05-2016") == date(2016, 5, 14)
    assert parse_date("14/05/2016") == date(2016, 5, 14)
    assert parse_date("2016-05-14") == date(2016, 5, 14)
    assert parse_date(datetime(2016, 5, 14, 0, 0)) == date(2016, 5, 14)
    assert parse_date("") is None


def test_clean_row_ok():
    r, errs, warns = clean_row(raw(roll_no=7.0, blood_group="b+ve", admission_no=412), TODAY)
    assert errs == []
    assert r["gender"] == "Female" and r["roll_no"] == 7 and r["blood_group"] == "B+"
    assert r["father_mobile"] == "919876543210" and r["father_name"] == "Father of Riya"
    assert r["admission_no"] == "412"


def test_clean_row_errors():
    r, errs, warns = clean_row(raw(first_name="", gender="x", date_of_birth="31-31-2016",
                                   father_mobile="12345", father_email="nope"), TODAY)
    joined = " | ".join(errs)
    for frag in ("First Name is required", "Gender 'x'", "Date of Birth", "Father Mobile", "Father Email"):
        assert frag in joined, frag


def test_future_dob_and_no_contact_warning():
    r, errs, warns = clean_row(raw(date_of_birth="01-01-2030", father_mobile=""), TODAY)
    assert any("past" in e for e in errs)
    assert any("won't get alerts" in w for w in warns)


def test_blood_group_variants():
    for given, want in (("B positive", "B+"), ("ab-", "AB-"), ("O +ve", "O+")):
        assert clean_row(raw(blood_group=given), TODAY)[0]["blood_group"] == want
    r, _, warns = clean_row(raw(blood_group="Z"), TODAY)
    assert r["blood_group"] == "" and warns


def test_cross_row_duplicates_and_blank_rows():
    rows = [raw(admission_no="A1", roll_no=1), {}, raw(admission_no="A1", roll_no=1, first_name="Asha")]
    out = check_rows(rows, TODAY)
    assert [c["row"] for c in out] == [2, 4]  # blank row 3 skipped
    errs = " ".join(out[1]["errors"])
    assert "Admission No A1 repeats row 2" in errs and "Roll No 1 repeats row 2" in errs


def test_group_name():
    assert group_name("Class 5", "A", "2026-27") == "Class 5 - A (2026-27)"


def test_template_roundtrip(tmp_path=None):
    content = build_template("Demo School", "2026-27")
    rows, missing = read_rows(content)
    assert missing == []
    assert len(rows) >= 1 and is_example_row(rows[0])  # the example row is recognised and can be skipped
    wb = load_workbook(__import__("io").BytesIO(content))
    ws = wb["Students"]
    assert ws.cell(row=1, column=1).value == "Class *"
    assert [c.value.rstrip(" *") for c in ws[1]] == [c[1] for c in COLUMNS]
    assert "How to fill" in wb.sheetnames


def test_read_rows_reports_missing_columns():
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["First Name", "Gender"])
    wb.active.append(["Riya", "F"])
    buf = BytesIO()
    wb.save(buf)
    rows, missing = read_rows(buf.getvalue())
    assert set(missing) == {"Class", "Section", "Date of Birth"}
    assert rows == [{"first_name": "Riya", "gender": "F"}]


def test_marathi_messages():
    ctx = localise_context({"entry_type": "Homework", "student_name": "रिया", "guardian_name": "सुरेश",
                            "title": "धडा ४", "subject": "मराठी", "due_date": "23 Oct", "body": "",
                            "portal_url": "x", "school_name": "डेमो शाळा"}, "mr")
    assert ctx["entry_type"] == "गृहपाठ"
    subject, body = email_template("Homework", "mr")
    assert render(subject, ctx) == "रिया साठी गृहपाठ: धडा ४"
    assert "{" not in render(body, ctx)
    assert email_template("Absence", "xx") == email_template("Absence", "en")
    assert template_params("Homework", ctx) == ["सुरेश", "रिया", "धडा ४", "23 Oct", "डेमो शाळा"]


def test_every_alert_type_has_templates():
    from schoolos.alerts.messages import DEFAULT_EMAIL, DEFAULT_EMAIL_MR, TEMPLATE_PARAMS

    assert set(TEMPLATE_PARAMS) == set(DEFAULT_EMAIL) == set(DEFAULT_EMAIL_MR)
